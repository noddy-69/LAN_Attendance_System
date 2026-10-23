"""
Report Tasks
============
Celery tasks that generate monthly reports, upload them to MinIO,
and return presigned GET URLs (valid 4 hours).

generate_attendance_csv(month, year, branch_id=None)
  daily_attendance JOIN employees → CSV → MinIO → presigned URL

generate_payroll_excel(month, year, branch_id=None)
  salary_deductions + daily_attendance → openpyxl Excel → MinIO → presigned URL

Both tasks use the synchronous psycopg2 driver (no asyncio in Celery workers)
and the synchronous minio-py SDK.

Monetary display: paise → rupees (÷ 100, rounded to 2 dp) for human-readable
output only.  Internal storage remains INTEGER paise throughout.
"""

from __future__ import annotations

import calendar
import csv
import io
import json
import logging
import uuid
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from typing import Optional

import openpyxl
from minio import Minio
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import settings
from app.tasks.celery_app import celery_app

log = logging.getLogger(__name__)

IST = timezone(timedelta(hours=5, minutes=30))

# ── Sync DB ────────────────────────────────────────────────────────────────────

_sync_factory: sessionmaker | None = None


def _get_factory() -> sessionmaker:
    global _sync_factory
    if _sync_factory is None:
        url = settings.DATABASE_URL.replace("postgresql+asyncpg://", "postgresql://")
        _sync_factory = sessionmaker(
            bind=create_engine(url, pool_size=3, pool_pre_ping=True),
            expire_on_commit=False,
        )
    return _sync_factory


# ── MinIO ──────────────────────────────────────────────────────────────────────

def _upload_and_presign(object_name: str, data: bytes, content_type: str) -> str:
    """Upload bytes to MinIO and return a 4-hour presigned GET URL."""
    client = Minio(
        settings.MINIO_ENDPOINT,
        access_key=settings.MINIO_ACCESS_KEY,
        secret_key=settings.MINIO_SECRET_KEY,
        secure=settings.MINIO_SECURE,
    )
    bucket = settings.MINIO_BUCKET_ATTENDANCE
    if not client.bucket_exists(bucket):
        client.make_bucket(bucket)

    client.put_object(
        bucket,
        object_name,
        io.BytesIO(data),
        length=len(data),
        content_type=content_type,
    )
    return client.presigned_get_object(
        bucket,
        object_name,
        expires=timedelta(hours=4),
    )


# ── Formatting helpers ─────────────────────────────────────────────────────────

def _ist_str(dt: datetime | None) -> str:
    if not dt:
        return ""
    return dt.astimezone(IST).strftime("%Y-%m-%d %H:%M")


def _work_hours(minutes: int | None) -> str:
    if minutes is None:
        return ""
    h, m = divmod(minutes, 60)
    return f"{h}h {m:02d}m"


def _hm(minutes: int | None) -> str:
    """Compact hours+minutes for durations like late/overtime (e.g. 90 -> '1h 30m')."""
    if not minutes:
        return "0m"
    h, m = divmod(int(minutes), 60)
    return f"{h}h {m:02d}m" if h else f"{m}m"


_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_CENTER = Alignment(horizontal="center", vertical="center")


def _style_headers(ws, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER


def _autofit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)


# ── Attendance data ────────────────────────────────────────────────────────────

_ATT_HEADERS = [
    "Employee Name",
    "Employee Code",
    "Date",
    "Check-In (IST)",
    "Check-Out (IST)",
    "Work Hours",
    "Status",
    "Is Late",
    "Late Time",
    "Overtime",
    "Source",
]


def _attendance_rows(
    session: Session,
    month: int,
    year: int,
    branch_id: uuid.UUID | None,
) -> list[list]:
    from app.models.daily_attendance import DailyAttendance
    from app.models.employee import Employee
    from app.models.role import Role

    first_day = date(year, month, 1)
    last_day = date(year, month, monthrange(year, month)[1])

    stmt = (
        select(DailyAttendance, Employee)
        .join(Employee, DailyAttendance.employee_id == Employee.id)
        .where(
            DailyAttendance.date >= first_day,
            DailyAttendance.date <= last_day,
            DailyAttendance.deleted_at.is_(None),
            Employee.deleted_at.is_(None),
            ~Employee.role.has(Role.name == "service"),
        )
    )
    if branch_id:
        stmt = stmt.where(Employee.branch_id == branch_id)
    stmt = stmt.order_by(DailyAttendance.date.desc(), Employee.employee_code)

    rows = []
    for att, emp in session.execute(stmt):
        rows.append([
            f"{emp.first_name} {emp.last_name}",
            emp.employee_code,
            att.date.isoformat(),
            _ist_str(att.check_in),
            _ist_str(att.check_out),
            _work_hours(att.work_minutes),
            att.status.value,
            "Yes" if att.is_late else "No",
            _hm(att.late_minutes),
            _hm(att.overtime_minutes),
            att.source.value,
        ])
    return rows


# ── Payroll data ───────────────────────────────────────────────────────────────

_PAYROLL_HEADERS = [
    "Employee Name",
    "Employee Code",
    "Month",
    "Year",
    "Days Present",
    "Days Absent",
    "Days on Leave",
    "Late Count",
    "Penalty Units",
    "Net Penalty Days",
    "Monthly Salary (Rs)",
    "Daily Salary (Rs)",
    "Total Deduction (Rs)",
]


def _payroll_rows(
    session: Session,
    month: int,
    year: int,
    branch_id: uuid.UUID | None,
) -> list[list]:
    from app.models.daily_attendance import DailyAttendance, AttendanceStatus
    from app.models.employee import Employee
    from app.models.payroll import SalaryDeduction
    from app.models.role import Role
    from app.models.system import SystemPolicy
    from app.utils.payroll_calc import DeductionPolicy, compute_deduction

    def _policy_raw(key: str):
        return session.execute(
            select(SystemPolicy.value).where(
                SystemPolicy.key == key,
                SystemPolicy.deleted_at.is_(None),
            )
        ).scalar_one_or_none()

    def _policy_int(key: str, default: int) -> int:
        row = _policy_raw(key)
        try:
            return int(json.loads(row)) if row is not None else default
        except (TypeError, ValueError):
            return default

    def _policy_float(key: str, default: float) -> float:
        row = _policy_raw(key)
        try:
            return float(json.loads(row)) if row is not None else default
        except (TypeError, ValueError):
            return default

    policy = DeductionPolicy(
        late_days_per_unit=_policy_int("payroll.late_days_per_unit", settings.DEFAULT_LATE_DAYS_PER_UNIT),
        penalty_percent_per_unit=_policy_int("payroll.penalty_percent_per_unit", settings.DEFAULT_PENALTY_PERCENT_PER_UNIT),
        grace_days=_policy_float("payroll.grace_days", settings.DEFAULT_GRACE_DAYS),
        salary_days_per_month=_policy_int("payroll.salary_days_per_month", settings.DEFAULT_SALARY_DAYS_PER_MONTH),
    )

    first_day = date(year, month, 1)
    last_day = date(year, month, monthrange(year, month)[1])
    month_name = calendar.month_name[month]

    present_statuses = {
        AttendanceStatus.PRESENT.value,
        AttendanceStatus.LATE.value,
        AttendanceStatus.HALF_DAY.value,
    }

    emp_stmt = (
        select(Employee)
        .where(
            Employee.is_active.is_(True),
            Employee.deleted_at.is_(None),
            ~Employee.role.has(Role.name == "service"),
        )
        .order_by(Employee.employee_code)
    )
    if branch_id:
        emp_stmt = emp_stmt.where(Employee.branch_id == branch_id)
    employees = session.execute(emp_stmt).scalars().all()

    ded_map: dict[uuid.UUID, SalaryDeduction] = {
        d.employee_id: d
        for d in session.execute(
            select(SalaryDeduction).where(
                SalaryDeduction.month == month,
                SalaryDeduction.year == year,
                SalaryDeduction.deleted_at.is_(None),
            )
        ).scalars().all()
    }

    rows = []
    for emp in employees:
        att_rows = session.execute(
            select(DailyAttendance.status, DailyAttendance.is_late).where(
                DailyAttendance.employee_id == emp.id,
                DailyAttendance.date >= first_day,
                DailyAttendance.date <= last_day,
                DailyAttendance.deleted_at.is_(None),
            )
        ).all()

        days_present  = sum(1 for r in att_rows if r.status.value in present_statuses)
        days_absent   = sum(1 for r in att_rows if r.status == AttendanceStatus.ABSENT)
        days_on_leave = sum(1 for r in att_rows if r.status == AttendanceStatus.ON_LEAVE)

        ded = ded_map.get(emp.id)
        if ded and ded.daily_salary_paise is not None:
            late_count    = ded.late_count
            units         = ded.penalty_units
            net_x100      = ded.net_penalty_days_x100
            monthly_paise = ded.monthly_salary_paise or emp.monthly_salary_paise or 0
            daily_paise   = ded.daily_salary_paise
            total_paise   = ded.total_deduction_paise
        else:
            late_count    = sum(1 for r in att_rows if r.is_late)
            monthly_paise = emp.monthly_salary_paise or 0
            calc          = compute_deduction(late_count, monthly_paise, policy)
            units         = calc["penalty_units"]
            net_x100      = calc["net_penalty_days_x100"]
            daily_paise   = calc["daily_salary_paise"]
            total_paise   = calc["total_deduction_paise"]

        rows.append([
            f"{emp.first_name} {emp.last_name}",
            emp.employee_code,
            month_name,
            year,
            days_present,
            days_absent,
            days_on_leave,
            late_count,
            units,
            round(net_x100 / 100, 2),
            round(monthly_paise / 100, 2),  # paise → ₹ for display
            round(daily_paise / 100, 2),
            round(total_paise / 100, 2),
        ])
    return rows


# ── Celery tasks ───────────────────────────────────────────────────────────────

@celery_app.task(
    name="app.tasks.reports.generate_attendance_csv",
    acks_late=True,
)
def generate_attendance_csv(
    month: int,
    year: int,
    branch_id: Optional[str] = None,
) -> dict:
    """
    Generate a monthly attendance CSV, upload to MinIO, return a 4-hour presigned URL.

    Parameters
    ----------
    month     : 1–12
    year      : e.g. 2026
    branch_id : UUID string or None (all branches)
    """
    bid = uuid.UUID(branch_id) if branch_id else None

    with _get_factory()() as session:
        rows = _attendance_rows(session, month, year, bid)

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_ATT_HEADERS)
    writer.writerows(rows)
    csv_bytes = buf.getvalue().encode("utf-8-sig")   # BOM for Excel compatibility

    object_name = f"reports/attendance_{year}_{month:02d}_{branch_id or 'all'}.csv"
    url = _upload_and_presign(object_name, csv_bytes, "text/csv; charset=utf-8")

    log.info(
        "generate_attendance_csv: %s/%d rows=%d uploaded=%s",
        calendar.month_name[month], year, len(rows), object_name,
    )
    return {"url": url, "rows": len(rows), "month": month, "year": year}


@celery_app.task(
    name="app.tasks.reports.generate_payroll_excel",
    acks_late=True,
)
def generate_payroll_excel(
    month: int,
    year: int,
    branch_id: Optional[str] = None,
) -> dict:
    """
    Generate a monthly payroll Excel report, upload to MinIO, return a presigned URL.

    Parameters
    ----------
    month     : 1–12
    year      : e.g. 2026
    branch_id : UUID string or None (all branches)
    """
    bid = uuid.UUID(branch_id) if branch_id else None

    with _get_factory()() as session:
        rows = _payroll_rows(session, month, year, bid)

    month_name = calendar.month_name[month]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{month_name[:3]} {year}"

    _style_headers(ws, _PAYROLL_HEADERS)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    _autofit(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    excel_bytes = buf.read()

    object_name = f"reports/payroll_{year}_{month:02d}_{branch_id or 'all'}.xlsx"
    url = _upload_and_presign(
        object_name,
        excel_bytes,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    log.info(
        "generate_payroll_excel: %s/%d rows=%d uploaded=%s",
        month_name, year, len(rows), object_name,
    )
    return {"url": url, "rows": len(rows), "month": month, "year": year}
