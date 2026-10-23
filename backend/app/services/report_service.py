"""
Report Service
==============
Generates downloadable attendance and leave reports as Excel (.xlsx) or CSV.

IST note: all DB timestamps are UTC; this service converts to IST (UTC+5:30)
for display in exported files.  This is the only place UTC→IST conversion
happens server-side — all other IST display is a frontend concern.
"""
import calendar
import csv
import io
import uuid
from datetime import date, datetime, timedelta, timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.daily_attendance import AttendanceStatus, DailyAttendance
from app.models.employee import Employee
from app.models.leave import LeaveBalance, LeaveRequest, LeaveStatus, LeaveType

IST = timezone(timedelta(hours=5, minutes=30))

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_CENTER = Alignment(horizontal="center", vertical="center")


def _ist_str(dt: datetime | None) -> str:
    """Format a UTC datetime as IST string, or empty string if None."""
    if not dt:
        return ""
    return dt.astimezone(IST).strftime("%Y-%m-%d %H:%M")


def _work_hours(minutes: int | None) -> str:
    if minutes is None:
        return ""
    h, m = divmod(minutes, 60)
    return f"{h}h {m:02d}m"


def _style_header_row(ws, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER


def _autofit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)


# ── Attendance report ─────────────────────────────────────────────────────────

_ATT_HEADERS = [
    "Employee Code",
    "Employee Name",
    "Date",
    "Check In",
    "Check Out",
    "Work Hours",
    "Late",
    "Late Minutes",
    "Overtime",
    "Status",
    "Source",
]


async def _attendance_rows(
    db: AsyncSession,
    from_date: date | None,
    to_date: date | None,
    employee_id: uuid.UUID | None,
    branch_id: uuid.UUID | None,
    status: AttendanceStatus | None,
) -> list[list]:
    stmt = (
        select(DailyAttendance)
        .options(selectinload(DailyAttendance.employee))
        .where(DailyAttendance.deleted_at.is_(None))
    )
    if employee_id:
        stmt = stmt.where(DailyAttendance.employee_id == employee_id)
    if from_date:
        stmt = stmt.where(DailyAttendance.date >= from_date)
    if to_date:
        stmt = stmt.where(DailyAttendance.date <= to_date)
    if status:
        stmt = stmt.where(DailyAttendance.status == status)
    if branch_id:
        stmt = stmt.join(Employee).where(
            Employee.branch_id == branch_id,
            Employee.deleted_at.is_(None),
        )
        stmt = stmt.order_by(DailyAttendance.date.desc(), Employee.employee_code)
    else:
        stmt = stmt.order_by(DailyAttendance.date.desc())

    result = await db.execute(stmt)
    records = list(result.scalars().all())

    rows = []
    for r in records:
        emp = r.employee
        rows.append([
            emp.employee_code,
            f"{emp.first_name} {emp.last_name}",
            r.date.isoformat(),
            _ist_str(r.check_in),
            _ist_str(r.check_out),
            _work_hours(r.work_minutes),
            "Yes" if r.is_late else "No",
            r.late_minutes,
            r.overtime_minutes,
            r.status.value,
            r.source.value,
        ])
    return rows


async def generate_attendance_excel(
    db: AsyncSession,
    from_date: date | None,
    to_date: date | None,
    employee_id: uuid.UUID | None = None,
    branch_id: uuid.UUID | None = None,
    status: AttendanceStatus | None = None,
) -> bytes:
    rows = await _attendance_rows(db, from_date, to_date, employee_id, branch_id, status)

    wb = openpyxl.Workbook()
    ws = wb.active
    date_range = f"{from_date or 'all'}_to_{to_date or 'all'}"
    ws.title = f"Attendance {date_range}"[:31]

    _style_header_row(ws, _ATT_HEADERS)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    _autofit(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def generate_attendance_csv(
    db: AsyncSession,
    from_date: date | None,
    to_date: date | None,
    employee_id: uuid.UUID | None = None,
    branch_id: uuid.UUID | None = None,
    status: AttendanceStatus | None = None,
) -> bytes:
    rows = await _attendance_rows(db, from_date, to_date, employee_id, branch_id, status)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_ATT_HEADERS)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


# ── Leave report ──────────────────────────────────────────────────────────────

_LEAVE_HEADERS = [
    "Employee Name",
    "Employee Code",
    "Leave Type",
    "Leave Code",
    "Start Date",
    "End Date",
    "Days",
    "Status",
    "Reason",
    "Manager",
    "Manager Action Date",
    "Manager Remarks",
    "HR",
    "HR Action Date",
    "HR Remarks",
    "Applied At",
]


async def _leave_rows(
    db: AsyncSession,
    from_date: date | None,
    to_date: date | None,
    employee_id: uuid.UUID | None,
    status: LeaveStatus | None,
) -> list[list]:
    _LEAVE_EAGER = [
        selectinload(LeaveRequest.employee),
        selectinload(LeaveRequest.leave_type),
        selectinload(LeaveRequest.approving_manager),
        selectinload(LeaveRequest.approving_hr),
    ]
    stmt = (
        select(LeaveRequest)
        .options(*_LEAVE_EAGER)
        .where(LeaveRequest.deleted_at.is_(None))
    )
    if employee_id:
        stmt = stmt.where(LeaveRequest.employee_id == employee_id)
    if from_date:
        stmt = stmt.where(LeaveRequest.start_date >= from_date)
    if to_date:
        stmt = stmt.where(LeaveRequest.end_date <= to_date)
    if status:
        stmt = stmt.where(LeaveRequest.status == status)
    stmt = stmt.order_by(LeaveRequest.created_at.desc())

    result = await db.execute(stmt)
    records = list(result.scalars().all())

    rows = []
    for r in records:
        emp = r.employee
        mgr = r.approving_manager
        hr  = r.approving_hr
        rows.append([
            f"{emp.first_name} {emp.last_name}",
            emp.employee_code,
            r.leave_type.name,
            r.leave_type.code,
            r.start_date.isoformat(),
            r.end_date.isoformat(),
            r.days_requested,
            r.status.value,
            r.reason,
            f"{mgr.first_name} {mgr.last_name}" if mgr else "",
            _ist_str(r.manager_action_at),
            r.manager_remarks or "",
            f"{hr.first_name} {hr.last_name}" if hr else "",
            _ist_str(r.hr_action_at),
            r.hr_remarks or "",
            _ist_str(r.created_at),
        ])
    return rows


async def generate_leave_excel(
    db: AsyncSession,
    from_date: date | None,
    to_date: date | None,
    employee_id: uuid.UUID | None = None,
    status: LeaveStatus | None = None,
) -> bytes:
    rows = await _leave_rows(db, from_date, to_date, employee_id, status)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leave Report"

    _style_header_row(ws, _LEAVE_HEADERS)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    _autofit(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def generate_leave_csv(
    db: AsyncSession,
    from_date: date | None,
    to_date: date | None,
    employee_id: uuid.UUID | None = None,
    status: LeaveStatus | None = None,
) -> bytes:
    rows = await _leave_rows(db, from_date, to_date, employee_id, status)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_LEAVE_HEADERS)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")
